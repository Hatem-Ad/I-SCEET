"""
================================================================================
 I-SCEET — TM Builder v1.0
 File: orchestrator/tm_builder.py
================================================================================
 Parses M6 v2 raw text (Sections 1/2/3/3.5/4, multi-link aware) into
 TM-xxx.xlsx.
================================================================================
"""

import os
import re

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    raise ImportError("openpyxl required: pip install openpyxl")

HEADER_BG = "1F3864"
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Arial", size=10)
BODY_FONT = Font(name="Arial", size=9.5)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
BODY_ALIGN = Alignment(vertical="top", wrap_text=True)

STATUS_COLORS = {"COMPLETE": "E2EFDA", "PARTIAL": "FFF2CC", "GAP": "FCE4D6"}


def _strip_think_tags(text: str) -> str:
    if not text:
        return text
    text = re.sub(r'<think>.*?</think>', '', text, flags=re.DOTALL)
    return re.sub(r'</?think>', '', text).strip()


def _write_header_row(ws, headers):
    for j, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=j, value=h)
        cell.font = HEADER_FONT
        cell.fill = PatternFill("solid", fgColor=HEADER_BG)
        cell.alignment = HEADER_ALIGN


def _parse_pipe_table(block: str) -> list:
    rows = []
    for line in block.split('\n'):
        line = line.strip()
        if not line.startswith('|'):
            continue
        cells = [c.strip() for c in line.strip('|').split('|')]
        if all(re.fullmatch(r'-+', c) for c in cells):
            continue
        rows.append(cells)
    return rows


def _extract_section(text: str, section_marker: str, next_markers: list) -> str:
    next_pattern = '|'.join(re.escape(m) for m in next_markers)
    m = re.search(
        rf'{re.escape(section_marker)}(.*?)(?=(?:{next_pattern})|\Z)',
        text, re.DOTALL | re.IGNORECASE
    )
    return m.group(1) if m else ''


def _write_table_sheet(wb, sheet_name, rows, status_col_idx=None):
    ws = wb.create_sheet(sheet_name)
    if not rows:
        ws.cell(row=1, column=1, value=f"[{sheet_name} not found or could not be parsed]")
        return ws
    _write_header_row(ws, rows[0])
    for i, row in enumerate(rows[1:], 2):
        for j, val in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=val)
            c.font = BODY_FONT
            c.alignment = BODY_ALIGN
            if status_col_idx and j == status_col_idx and val in STATUS_COLORS:
                c.fill = PatternFill("solid", fgColor=STATUS_COLORS[val])
            elif i % 2 == 0:
                c.fill = PatternFill("solid", fgColor="F2F2F2")
    for j in range(1, len(rows[0]) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 18
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def _parse_gaps(block: str) -> list:
    items = []
    for line in block.split('\n'):
        line = line.strip().lstrip('-').strip()
        if line.startswith('**') and '**' in line[2:]:
            m = re.match(r'\*\*(.+?)\*\*:?\s*(.*)', line)
            if m:
                items.append((m.group(1).strip(), m.group(2).strip()))
        elif ':' in line and line:
            k, _, v = line.partition(':')
            items.append((k.strip(), v.strip()))
    return items


def build_tm(m6_output: str, output_path: str, version: str = "1.0") -> str:
    """Build TM-xxx.xlsx from M6 v2 raw text (Sections 1/2/3/3.5/4)."""
    text = _strip_think_tags(m6_output)

    section_markers = [
        "SECTION 1", "SECTION 2", "SECTION 3.5", "SECTION 3", "SECTION 4"
    ]
    sec1 = _extract_section(text, "SECTION 1", ["SECTION 2"])
    sec2 = _extract_section(text, "SECTION 2", ["SECTION 3.5", "SECTION 3"])
    sec3 = _extract_section(text, "SECTION 3:", ["SECTION 3.5", "SECTION 4"])
    sec35 = _extract_section(text, "SECTION 3.5", ["SECTION 4"])
    sec4 = _extract_section(text, "SECTION 4", ["CERTIFICATION READINESS"])

    cert_match = re.search(r'CERTIFICATION READINESS:\s*(.+)', text)
    cert_readiness = cert_match.group(1).strip() if cert_match else "UNKNOWN"

    gaps = _parse_gaps(sec3)
    meaningful_gaps = [
        (cat, detail) for cat, detail in gaps
        if (detail or cat).strip().lower() not in {
            "none", "no gap", "no gaps", "n/a", "na", "complete", "all complete"
        }
    ]
    if meaningful_gaps and "NOT READY" not in cert_readiness.upper():
        cert_readiness = "NOT READY — Traceability gaps detected"

    wb = Workbook()
    ws_info = wb.active
    ws_info.title = "INFO"
    for i, (k, v) in enumerate([
        ("Document", "Traceability Matrix"), ("Version", version),
        ("Certification Readiness", cert_readiness),
    ], 1):
        ws_info.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws_info.cell(row=i, column=2, value=v)
    ws_info.column_dimensions["A"].width = 24
    ws_info.column_dimensions["B"].width = 50

    fwd_rows = _parse_pipe_table(sec1)
    status_idx = len(fwd_rows[0]) if fwd_rows else None
    _write_table_sheet(wb, "Forward Traceability", fwd_rows, status_col_idx=status_idx)

    bwd_rows = _parse_pipe_table(sec2)
    status_idx2 = len(bwd_rows[0]) if bwd_rows else None
    _write_table_sheet(wb, "Backward Traceability", bwd_rows, status_col_idx=status_idx2)

    agg_rows = _parse_pipe_table(sec35)
    status_idx3 = len(agg_rows[0]) if agg_rows else None
    _write_table_sheet(wb, "Aggregated Coverage", agg_rows, status_col_idx=status_idx3)

    ws_gaps = wb.create_sheet("Traceability Gaps")
    _write_header_row(ws_gaps, ["Gap Category", "Details"])
    if gaps:
        for i, (cat, detail) in enumerate(gaps, 2):
            ws_gaps.cell(row=i, column=1, value=cat).font = BODY_FONT
            c2 = ws_gaps.cell(row=i, column=2, value=detail or "None")
            c2.font = BODY_FONT
            c2.alignment = BODY_ALIGN
    ws_gaps.column_dimensions["A"].width = 35
    ws_gaps.column_dimensions["B"].width = 60

    ws_cov = wb.create_sheet("Coverage Summary")
    row_i = 1
    for line in [l.strip() for l in sec4.split('\n') if l.strip()]:
        line = line.lstrip('-').strip().replace('**', '')
        if not line or line.upper() == 'COVERAGE SUMMARY':
            continue
        if ':' in line:
            k, _, v = line.partition(':')
            ws_cov.cell(row=row_i, column=1, value=k.strip()).font = Font(bold=True)
            ws_cov.cell(row=row_i, column=2, value=v.strip())
            row_i += 1
    ws_cov.cell(row=row_i + 1, column=1,
               value="Certification Readiness").font = Font(bold=True)
    ws_cov.cell(row=row_i + 1, column=2, value=cert_readiness)
    ws_cov.column_dimensions["A"].width = 30
    ws_cov.column_dimensions["B"].width = 40

    os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
    wb.save(output_path)
    return output_path
