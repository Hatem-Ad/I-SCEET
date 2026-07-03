"""
================================================================================
 I-SCEET — DD Builder v1.0
 File: orchestrator/dd_builder.py
================================================================================
 Parses [VAR]/[CONST]/[TYPE] Data Dictionary tags from M2 raw text ->
 generates DD-xxx.xlsx (3 sheets: Variables, Constants, Types).

 Reuses sddd_builder._extract_dd_tags so the DD and the SDDD docx are
 always parsed from the exact same tag syntax — no risk of divergence
 between the two documents.
================================================================================
"""

import os

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    raise ImportError("openpyxl required: pip install openpyxl")

from sddd_builder import _extract_dd_tags

HEADER_BG = "1F3864"
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Arial", size=10)
BODY_FONT = Font(name="Arial", size=9.5)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
BODY_ALIGN = Alignment(vertical="top", wrap_text=True)

VARIABLES_COLUMNS = [
    ("Name", "name"), ("Type", "type"), ("Range", "range"),
    ("Description", "description"), ("Unit", "unit"),
    ("Initializer", "initializer"), ("Producer(s)", "producers"),
    ("Consumer(s)", "consumers"), ("Comments", "comments"),
]
CONSTANTS_COLUMNS = [
    ("Name", "name"), ("Type", "type"), ("Definition", "definition"),
    ("Description", "description"), ("Comments", "comments"),
]
TYPES_COLUMNS = [
    ("Name", "name"), ("Type", "type"), ("Definition", "definition"),
    ("Description", "description"), ("Comments", "comments"),
]


def _write_sheet(ws, columns, rows):
    for j, (header, _) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=j, value=header)
        cell.font = HEADER_FONT
        cell.fill = PatternFill("solid", fgColor=HEADER_BG)
        cell.alignment = HEADER_ALIGN

    for i, row in enumerate(rows, start=2):
        for j, (_, key) in enumerate(columns, 1):
            cell = ws.cell(row=i, column=j, value=row.get(key, '') or '-')
            cell.font = BODY_FONT
            cell.alignment = BODY_ALIGN
            if i % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F2F2F2")

    widths = {1: 22, 2: 16}
    for j in range(1, len(columns) + 1):
        letter = get_column_letter(j)
        ws.column_dimensions[letter].width = widths.get(j, 28)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def build_dd(m2_output: str, output_path: str, version: str = "1.0") -> str:
    """Build DD-xxx.xlsx from M2 raw text output (same text sddd_builder
    consumes — the [VAR]/[CONST]/[TYPE] tags are shared ground truth)."""
    dd = _extract_dd_tags(m2_output)

    wb = Workbook()
    ws_var = wb.active
    ws_var.title = "Variables"
    _write_sheet(ws_var, VARIABLES_COLUMNS, dd['variables'])

    ws_const = wb.create_sheet("Constants")
    _write_sheet(ws_const, CONSTANTS_COLUMNS, dd['constants'])

    ws_types = wb.create_sheet("Types")
    _write_sheet(ws_types, TYPES_COLUMNS, dd['types'])

    os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
    wb.save(output_path)
    return output_path
