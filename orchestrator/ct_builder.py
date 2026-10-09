"""
================================================================================
 I-SCEET — CT Builder v1.0
 File: orchestrator/ct_builder.py
 Copyright (c) 2026 Hatem Adouani — MIT License
================================================================================
 Derives CT-xxx.xlsx (2 sheets: Call Tree, Activation Scenario) from each
 function's "Called Functions" field — already reliably parsed by
 sddd_builder.py. No new M2 tag was needed for this.

 LIMITATION (flagged honestly, not hidden): the raw "Called Functions"
 field does not carry *conditional* information (whether the call happens
 inside a specific Cause/Effect branch or unconditionally). The "Condition"
 column is therefore populated as "-" (unconditional) for now. A future
 iteration could cross-reference each LLR's Causes/Effects "Effect" text
 against the Called Functions list to infer conditional calls, but that
 needs a more consistent "CALL <function>" phrasing in M2's Implementation
 lines than is currently guaranteed — left as a follow-up, not guessed at.
================================================================================
"""

import os
import re

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    raise ImportError("openpyxl required: pip install openpyxl")

try:
    from orchestrator.sddd_builder import _parse_m2_output
except ImportError:
    from sddd_builder import _parse_m2_output  # fallback: run standalone from orchestrator/

HEADER_BG = "1F3864"
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Arial", size=10)
BODY_FONT = Font(name="Arial", size=9.5)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
BODY_ALIGN = Alignment(vertical="top", wrap_text=True)

MAX_DEPTH = 7
LEVEL_LABELS = ["1st Level", "2nd Level", "3rd Level", "4th Level",
                "5th Level", "6th Level", "7th Level"]


def _parse_call_map(m2_output: str) -> dict:
    """function name -> list of called function names, from each
    function's 'Called Functions' field."""
    parsed = _parse_m2_output(m2_output)
    call_map = {}
    for fn in parsed.get('functions', []):
        name = fn.get('Name', '').strip()
        if not name:
            continue
        raw = fn.get('Called Functions', '')
        called = []
        if raw and raw.strip().lower() not in ('none', '-', ''):
            for part in re.split(r',', raw):
                cleaned = re.sub(r'[`()]', '', part).strip()
                if cleaned:
                    called.append(cleaned)
        call_map[name] = called
    return call_map


def _detect_roots(call_map: dict) -> list:
    """1st-level functions = never called by any other function in this
    subsystem (auto-detected, not guessed)."""
    called_anywhere = {c for calls in call_map.values() for c in calls}
    roots = [fn for fn in call_map if fn not in called_anywhere]
    return roots if roots else list(call_map.keys())


def _write_scenario_node(ws, call_map, node, depth, row_i, visited):
    """Pre-order DFS: write `node` alone in its depth column, then each
    child immediately follows (before moving to the next sibling) — an
    outline layout, not one row per full path."""
    col = depth + 1
    cell = ws.cell(row=row_i, column=col, value=node)
    cell.font = BODY_FONT
    cell.alignment = BODY_ALIGN
    row_i += 1

    if depth + 1 >= MAX_DEPTH:
        return row_i

    for child in call_map.get(node, []):
        if child in visited:
            cell = ws.cell(row=row_i, column=col + 1, value=f"{child} (cycle — truncated)")
            cell.font = BODY_FONT
            cell.alignment = BODY_ALIGN
            row_i += 1
            continue
        row_i = _write_scenario_node(ws, call_map, child, depth + 1, row_i, visited | {child})

    return row_i


def _write_activation_scenario(ws, call_map: dict, roots: list):
    row_i = 2
    for root in roots:
        title_cell = ws.cell(row=row_i, column=1, value=f"Activation scenario of ({root})")
        title_cell.font = Font(bold=True, name="Arial", size=10)
        ws.merge_cells(start_row=row_i, start_column=1, end_row=row_i, end_column=MAX_DEPTH)
        row_i += 1

        row_i = _write_scenario_node(ws, call_map, root, 0, row_i, {root})
        row_i += 1  # blank separator row before the next root's block

    return row_i


def _write_header(ws, headers):
    for j, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=j, value=h)
        cell.font = HEADER_FONT
        cell.fill = PatternFill("solid", fgColor=HEADER_BG)
        cell.alignment = HEADER_ALIGN


def build_ct(m2_output: str, output_path: str, version: str = "1.0") -> str:
    """Build CT-xxx.xlsx from M2 raw text output."""
    call_map = _parse_call_map(m2_output)

    wb = Workbook()

    # ── Sheet 1: Call Tree (Calling function | Called function | Condition)
    ws_tree = wb.active
    ws_tree.title = "Call Tree"
    _write_header(ws_tree, ["Calling Function", "Called Function", "Condition"])
    row_i = 2
    for caller, callees in call_map.items():
        if not callees:
            continue
        for callee in callees:
            for j, val in enumerate([caller, callee, "-"], 1):
                cell = ws_tree.cell(row=row_i, column=j, value=val)
                cell.font = BODY_FONT
                cell.alignment = BODY_ALIGN
                if row_i % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F2F2F2")
            row_i += 1
    for j, w in enumerate([28, 28, 14], 1):
        ws_tree.column_dimensions[get_column_letter(j)].width = w
    ws_tree.freeze_panes = "A2"
    if row_i > 2:
        ws_tree.auto_filter.ref = ws_tree.dimensions

    # ── Sheet 2: Activation Scenario (1st..7th Level, outline layout)
    ws_scenario = wb.create_sheet("Activation Scenario")
    _write_header(ws_scenario, LEVEL_LABELS)
    roots = _detect_roots(call_map)
    _write_activation_scenario(ws_scenario, call_map, roots)
    for j in range(1, MAX_DEPTH + 1):
        ws_scenario.column_dimensions[get_column_letter(j)].width = 22
    ws_scenario.freeze_panes = "A2"

    os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
    wb.save(output_path)
    return output_path