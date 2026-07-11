"""
================================================================================
 I-SCEET — SUTC Builder v1.0
 File: orchestrator/sutc_builder.py
================================================================================
 Parses M4 v3 raw text (5-sheet structure: A0/A1/A2 global, B0/B1 repeat
 per function) into SUTC-xxx.xlsx. Each function gets its own B0+B1 sheet
 pair, B1 named after the function (per SVS / M4 prompt v3).
================================================================================
"""

import os
import re

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    raise ImportError("openpyxl required: pip install openpyxl")

HEADER_BG = "1F3864"
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Arial", size=10)
BODY_FONT = Font(name="Arial", size=9.5)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
BODY_ALIGN = Alignment(vertical="top", wrap_text=True)


def _strip_think_tags(text: str) -> str:
    if not text:
        return text
    text = re.sub(r'<think>.*?</think>', '', text, flags=re.DOTALL)
    return re.sub(r'</?think>', '', text).strip()


def _clean_markdown_inline(text: str) -> str:
    if not text:
        return text
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    return re.sub(r'`([^`]+)`', r'\1', text)


def _sanitize_sheet_name(name: str, prefix: str = "") -> str:
    """Excel sheet names: max 31 chars, no []:*?/\\ characters."""
    name = re.sub(r'[\[\]:*?/\\]', '', name)
    full = f"{prefix}{name}" if prefix else name
    return full[:31]


def _write_header_row(ws, headers):
    for j, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=j, value=h)
        cell.font = HEADER_FONT
        cell.fill = PatternFill("solid", fgColor=HEADER_BG)
        cell.alignment = HEADER_ALIGN


def _write_kv_sheet(ws, pairs):
    """Simple two-column key/value sheet for A0/A2-style blocks."""
    for i, (k, v) in enumerate(pairs, 1):
        c1 = ws.cell(row=i, column=1, value=k)
        c1.font = Font(bold=True, name="Arial", size=9.5)
        c2 = ws.cell(row=i, column=2, value=v)
        c2.font = BODY_FONT
        c2.alignment = BODY_ALIGN
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 70


def _parse_labeled_block(text: str, label: str) -> str:
    """Extract the block of text after 'label:' up to the next top-level
    'SHEET'/'###' marker or end of string."""
    m = re.search(
        rf'{re.escape(label)}:\s*\n(.*?)(?=\n##\s*SHEET|\n#{{2,3}}\s|\Z)',
        text, re.DOTALL
    )
    return m.group(1).strip() if m else ''


def _parse_pipe_table(block: str) -> list:
    rows = []
    for line in block.split('\n'):
        line = line.strip()
        if not line.startswith('|'):
            continue
        cells = [c.strip() for c in line.strip('|').split('|')]
        if all(re.fullmatch(r'-+', c) for c in cells):
            continue
        rows.append(cells)
    return rows


def _parse_a0(text: str) -> dict:
    block = _parse_labeled_block(text, "A0")
    fields = {}
    for key in ["Source File", "CSC", "CSU", "Input Documents", "Functions Under Test"]:
        m = re.search(rf'{re.escape(key)}:\s*(.*?)(?=\n[A-Z][\w ]*:|\Z)', block, re.DOTALL)
        fields[key] = m.group(1).strip().replace('\n', ' ') if m else ''
    return fields


def _parse_a2(text: str) -> dict:
    block = _parse_labeled_block(text, "A2")
    fields = {}
    for key in ["Includes", "Variables", "Constants", "Types"]:
        m = re.search(
            rf'{re.escape(key)}:\s*(.*?)(?=\n[A-Z][\w ]*:|\n-{{5,}}|\n#{{2,3}}\s|\Z)',
            block, re.DOTALL
        )
        fields[key] = m.group(1).strip() if m else ''
    return fields


def _parse_functions(text: str) -> list:
    """Split on '### <function_name>() — LLR-ID(s)' headers."""
    pattern = re.compile(
        r'^#{2,3}\s+(.+?)\(\)\s*[—-]\s*(.+?)\n(.*?)(?=^#{2,3}\s+.+?\(\)|\Z)',
        re.DOTALL | re.MULTILINE
    )
    functions = []
    for m in pattern.finditer(text):
        name = m.group(1).strip()
        llr_ids = m.group(2).strip()
        body = m.group(3)
        functions.append({"name": name, "llr_ids": llr_ids, "body": body})
    return functions


def _parse_b0(body: str) -> list:
    b0_block = _parse_labeled_block(body, "B0 (Causes/Effects x UFT matrix)")
    if not b0_block:
        # fallback: grab first pipe table before "B1"
        m = re.search(r'(.*?)(?=\nB1)', body, re.DOTALL)
        b0_block = m.group(1) if m else body
    return _parse_pipe_table(b0_block)


def _parse_b1(body: str) -> list:
    """Parse UFT-ID: key: value blocks."""
    b1_match = re.search(r'B1[^\n]*:\s*\n(.*)', body, re.DOTALL)
    b1_block = b1_match.group(1) if b1_match else body
    cases = []
    for m in re.finditer(
        r'(SUTC_[A-Z0-9]+_UFT\d+):\s*\n(.*?)(?=\n(?:SUTC_[A-Z0-9]+_UFT\d+):|\Z)',
        b1_block, re.DOTALL
    ):
        uft_id, fields_text = m.group(1), m.group(2)
        fields = {"UFT-ID": uft_id}
        for key in ["Verifies", "Test-Type", "Stub", "Precondition (Init)",
                    "Input", "Expected", "Obtained", "Pass-Criteria", "Fail-Criteria"]:
            fm = re.search(rf'{re.escape(key)}:\s*(.*?)(?=\n\s*[A-Za-z][\w \-\(\)]*:|\Z)',
                            fields_text, re.DOTALL)
            fields[key] = _clean_markdown_inline(fm.group(1).strip()) if fm else ''
        cases.append(fields)
    return cases


def build_sutc(m4_output: str, output_path: str, version: str = "1.0") -> str:
    """Build SUTC-xxx.xlsx from M4 v3 raw text (5-sheet A0/A1/A2/B0/B1)."""
    text = _strip_think_tags(m4_output)

    wb = Workbook()
    a0 = _parse_a0(text)
    ws_a0 = wb.active
    ws_a0.title = "A0"
    _write_kv_sheet(ws_a0, list(a0.items()))

    a1_block = _parse_labeled_block(text, "A1")
    a1_rows = _parse_pipe_table(a1_block)
    ws_a1 = wb.create_sheet("A1")
    if a1_rows:
        _write_header_row(ws_a1, a1_rows[0])
        for i, row in enumerate(a1_rows[1:], 2):
            for j, val in enumerate(row, 1):
                c = ws_a1.cell(row=i, column=j, value=val)
                c.font = BODY_FONT
                c.alignment = BODY_ALIGN
    else:
        _write_header_row(ws_a1, ["Version", "Author", "Date", "Comment"])
        ws_a1.cell(row=2, column=1, value=version)

    a2 = _parse_a2(text)
    ws_a2 = wb.create_sheet("A2")
    _write_kv_sheet(ws_a2, list(a2.items()))

    functions = _parse_functions(text)
    total_uft = 0
    used_names = set()
    for fn in functions:
        b0_rows = _parse_b0(fn["body"])
        b1_cases = _parse_b1(fn["body"])
        total_uft += len(b1_cases)

        base_name = re.sub(r'\W+', '_', fn["name"])
        b0_name = _sanitize_sheet_name(base_name, "B0_")
        b1_name = _sanitize_sheet_name(base_name, "B1_")
        for nm in (b0_name, b1_name):
            suffix = 1
            orig = nm
            while nm in used_names:
                suffix += 1
                nm = f"{orig[:28]}_{suffix}"
            used_names.add(nm)

        ws_b0 = wb.create_sheet(b0_name)
        if b0_rows:
            _write_header_row(ws_b0, b0_rows[0])
            for i, row in enumerate(b0_rows[1:], 2):
                for j, val in enumerate(row, 1):
                    c = ws_b0.cell(row=i, column=j, value=val)
                    c.font = BODY_FONT
                    c.alignment = BODY_ALIGN
                    if i % 2 == 0:
                        c.fill = PatternFill("solid", fgColor="F2F2F2")
        else:
            _write_header_row(ws_b0, ["Row"])
            ws_b0.cell(row=2, column=1, value="[B0 matrix not found or could not be parsed]")

        ws_b1 = wb.create_sheet(b1_name)
        headers = ["UFT-ID", "Verifies", "Test-Type", "Stub", "Precondition (Init)",
                   "Input", "Expected", "Obtained", "Pass-Criteria", "Fail-Criteria"]
        _write_header_row(ws_b1, headers)
        for i, case in enumerate(b1_cases, 2):
            for j, key in enumerate(headers, 1):
                c = ws_b1.cell(row=i, column=j, value=case.get(key, ''))
                c.font = BODY_FONT
                c.alignment = BODY_ALIGN
                if i % 2 == 0:
                    c.fill = PatternFill("solid", fgColor="F2F2F2")
        for j in range(1, len(headers) + 1):
            ws_b1.column_dimensions[get_column_letter(j)].width = 20
        ws_b1.freeze_panes = "A2"

    # Coverage summary sheet
    ws_cov = wb.create_sheet("Coverage Summary")
    ws_cov.cell(row=1, column=1, value="Functions covered").font = Font(bold=True)
    ws_cov.cell(row=1, column=2, value=len(functions))
    ws_cov.cell(row=2, column=1, value="Total UFT cases generated").font = Font(bold=True)
    ws_cov.cell(row=2, column=2, value=total_uft)

    if not functions:
        ws_a0.cell(row=len(a0) + 2, column=1,
                  value="[No function blocks found or could not be parsed]")

    os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
    wb.save(output_path)
    return output_path
