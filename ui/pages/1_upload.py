"""
================================================================================
 I-SCEET — Page 1 : Upload documents + Nouveau projet
 File: ui/pages/1_upload.py
 Copyright (c) 2026 Hatem Adouani — MIT License
================================================================================
 v2: each file is saved immediately on upload (no all-or-nothing batch
 button) so a project can be built up incrementally - upload just the HW
 doc today, add the rest later, each new upload becomes the new version
 used by the next generation. Real text is extracted from docx/pdf/txt
 instead of a "[FILE: name]" placeholder.
================================================================================
"""

import streamlit as st
import sys, os

sys.path.append(os.path.join(os.path.dirname(__file__), "..", ".."))
sys.path.append(os.path.join(os.path.dirname(__file__), ".."))
from Backend.db import (create_project, get_all_projects, save_artifact,
                        get_artifact_by_module, get_artifact_history)
from theme import inject_theme, page_header, badge

UPLOAD_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "data", "uploads")
os.makedirs(UPLOAD_DIR, exist_ok=True)

st.set_page_config(page_title="Upload — I-SCEET", layout="wide")
inject_theme()
page_header("WORKSPACE 01 • INPUTS", "Programme d’entrée & configuration projet", "Create the active project and manage controlled source documents used by the generation pipeline.")


# ── TEXT EXTRACTION ────────────────────────────────────────────────────────────
def extract_text(file) -> str:
    """Extract real text content from an uploaded docx/pdf/txt file.
    Never returns a placeholder - callers should treat an empty result
    as extraction failure, not as 'no content'."""
    name = file.name.lower()
    file.seek(0)
    try:
        if name.endswith(".docx"):
            import docx
            d = docx.Document(file)
            parts = [p.text for p in d.paragraphs if p.text.strip()]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            return "\n".join(parts)
        elif name.endswith(".pdf"):
            file.seek(0)
            try:
                import pdfplumber
                parts = []
                with pdfplumber.open(file) as pdf:
                    for page in pdf.pages:
                        text = page.extract_text()
                        if text:
                            parts.append(text)
                return "\n".join(parts)
            except ImportError:
                pass  # fall through to pypdf/PyPDF2 below
            file.seek(0)
            try:
                from pypdf import PdfReader
            except ImportError:
                try:
                    from PyPDF2 import PdfReader
                except ImportError:
                    st.error(
                        "Aucune bibliothèque PDF installée on this server (tried "
                        "pdfplumber, pypdf, PyPDF2). Run "
                        "`pip install pdfplumber` (or `pypdf`) on the "
                        "Streamlit host and restart the app."
                    )
                    return ""
            reader = PdfReader(file)
            parts = [page.extract_text() or "" for page in reader.pages]
            return "\n".join(parts)
        elif name.endswith(".txt"):
            return file.read().decode("utf-8", errors="replace")
    except Exception as e:
        st.error(f"Échec d'extraction pour {file.name} : {e}")
        return ""
    return ""


# ── STEP 1 : PROJECT ──────────────────────────────────────────────────────────
st.markdown("### Étape 1 — Créer le projet")
col1, col2 = st.columns([2, 1])
with col1:
    project_name = st.text_input("Nom du projet", placeholder="ex: STM32F303_Bootloader_v1")
with col2:
    dal_level = st.selectbox("Niveau DAL", ["A", "B", "C", "D"])

if st.button("Créer le projet", type="primary", disabled=not project_name):
    existing = [p["name"] for p in get_all_projects()]
    if project_name in existing:
        st.error(f"Un projet nommé **{project_name}** existe déjà.")
    else:
        pid = create_project(project_name, dal_level)
        st.session_state["last_created_project_id"] = pid
        st.session_state["history_notice"] = (
            f"Projet **{project_name}** créé. Il n’est pas activé automatiquement : "
            "cliquez sur **Activer ce projet**."
        )
        st.switch_page("pages/6_historique.py")

st.divider()

active = st.session_state.get("active_project")
if not active:
    st.warning("Créez ou sélectionnez un projet ci-dessus avant d'uploader des documents.")
    st.stop()


# ── SHARED: one uploader that saves immediately ────────────────────────────────
def extract_version_from_docx(file):
    """Try to read the Version/Status from a docx cover-page table (the
    'Document ID'/'Version'/'Status' pattern used by every doc this
    project generates). Returns a version dict, or None if the file
    isn't a docx or doesn't have a parseable cover table - callers
    should fall back to normal auto-versioning (starts at 1.0) in that
    case, not assume 1.0 outright."""
    if not file.name.lower().endswith(".docx"):
        return None
    try:
        import docx
        file.seek(0)
        d = docx.Document(file)
        if not d.tables:
            return None
        fields = {row.cells[0].text.strip(): row.cells[1].text.strip()
                  for row in d.tables[0].rows if len(row.cells) >= 2}
        version_str = fields.get("Version", "").strip()
        status_str = fields.get("Status", "DRAFT").strip().upper()
        if not version_str:
            return None
        if "." in version_str:
            baseline_str, minor_str = version_str.split(".", 1)
        else:
            baseline_str, minor_str = version_str, "0"
        return {
            "baseline_number": int(baseline_str),
            "minor_number": int(minor_str),
            "status": "BASELINE" if "BASELINE" in status_str else "DRAFT",
        }
    except Exception:
        return None


def doc_uploader(doc_id, doc_label, module, col):
    """Render one uploader. Accepts one or more files (e.g. HW may have
    several reference manuals) — their extracted text is concatenated
    into a single artifact. Saves to DB immediately, no batch step.

    Guards against a real Streamlit gotcha: file_uploader keeps returning
    the same file(s) on every rerun until the widget is cleared, so a
    naive "if file: save" re-saves identical content on every rerun
    forever (this caused the reported 367x duplicate saves). A signature
    of the uploaded file set is kept in session_state and compared before
    saving — only a genuinely NEW file selection triggers a save."""
    with col:
        current = get_artifact_by_module(active["id"], module, doc_id)
        status_badge = badge(f"v{current['version']} \u2014 {current['created_at'][:16]}", "success") if current else badge("non t\u00e9l\u00e9vers\u00e9", "neutral")
        st.markdown(f"**{doc_id}** \u2014 {doc_label}", unsafe_allow_html=True)
        st.markdown(status_badge, unsafe_allow_html=True)

        history = get_artifact_history(active["id"], module, doc_id)
        if len(history) > 1:
            with st.expander(f"Historique ({len(history)} versions)"):
                for h in reversed(history):
                    marker = badge("baseline", "success") if h["status"] == "BASELINE" else badge("brouillon", "neutral")
                    st.markdown(
                        f"{marker} v{h['version']} \u2014 {h['created_at'][:16]}"
                        + (f" \u2014 {h['change_description']}" if h.get('change_description') else ""),
                        unsafe_allow_html=True
                    )
            st.caption(
                "T\u00e9l\u00e9verser un nouveau fichier ci-dessous cr\u00e9e une nouvelle version "
                "et remplace celle utilis\u00e9e pour la **prochaine** g\u00e9n\u00e9ration \u2014 "
                "les versions pr\u00e9c\u00e9dentes restent consultables ici, rien n'est perdu."
            )

        files = st.file_uploader(
            "Upload", type=["docx", "pdf", "txt"], key=f"up_{doc_id}",
            label_visibility="collapsed", accept_multiple_files=True
        )
        if not files:
            return

        signature = tuple((f.name, f.size) for f in files)
        sig_key = f"saved_sig_{doc_id}"
        if st.session_state.get(sig_key) == signature:
            return  # this exact selection already SUCCEEDED - stop the loop

        texts = []
        saved_paths = []
        detected_version = None
        for f in files:
            try:
                text = extract_text(f)
                if not text.strip():
                    st.error(f"Extraction impossible pour {f.name} — ignoré.")
                    continue
                texts.append(f"--- {f.name} ---\n{text}")
                path = os.path.join(UPLOAD_DIR, f"{active['id']}_{doc_id}_{f.name}")
                f.seek(0)
                with open(path, "wb") as out:
                    out.write(f.read())
                saved_paths.append(path)
                # Version detection only makes sense for our own docx
                # templates (cover-page Version field) - skip it outright
                # for non-docx files (PDFs, txt) rather than even calling
                # into that logic, and never let it block the save.
                if detected_version is None and f.name.lower().endswith(".docx"):
                    detected_version = extract_version_from_docx(f)
            except Exception as e:
                st.error(f"Erreur inattendue pour {f.name} : {e}")
                continue

        if not texts:
            return  # failed - do NOT mark signature, so the next rerun retries
                    # automatically once the underlying issue (e.g. missing
                    # PDF library) is fixed, instead of being silently stuck

        combined_text = "\n\n".join(texts)
        combined_path = saved_paths[0] if len(saved_paths) == 1 else ";".join(saved_paths)
        try:
            save_artifact(active["id"], module, doc_id, combined_text, combined_path, version=detected_version)
        except Exception as e:
            st.error(f"Échec de sauvegarde pour {doc_id} : {e}")
            return

        st.session_state[sig_key] = signature  # only mark AFTER confirmed success

        if detected_version:
            v = detected_version
            st.success(
                f"{doc_id} enregistré en v{v['baseline_number']}.{v['minor_number']} "
                f"({v['status']}) — {len(files)} file(s), version detected from the "
                f"document's own cover page."
            )
        else:
            st.success(f"{doc_id} enregistré ({len(files)} fichier(s), {len(combined_text):,} chars) — "
                       f"will be used in the next generation.")
        st.rerun()


# ── STEP 2 : PLANNING & STANDARDS ────────────────────────────────────────────
planning_documents = {
    "PSAC": "Plan for Software Aspects of Certification",
    "SDP":  "Software Development Plan",
    "SVP":  "Software Verification Plan",
    "SCMP": "Software Configuration Management Plan",
    "SQAP": "Software Quality Assurance Plan",
}

standard_documents = {
    "SRS":  "Software Requirements Standards",
    "SDS":  "Software Design Standards",
    "SCS":  "Software Code Standards",
    "SVS":  "Software Verification Standard",
}

planning_docs = {**planning_documents, **standard_documents}

st.markdown(f"### Étape 2 — Planning & Standards DO-178C ({len(planning_docs)} documents)")
st.caption(
    "Cette partie sert aussi à changer les documents de planning et de standards "
    "pour le prochain run. Chaque nouvel upload crée une nouvelle version en base "
    "et devient la référence utilisée par M1→M6 lors de la prochaine génération."
)

st.markdown("#### 2A — Planning documents")
cols = st.columns(2)
for i, (doc_id, doc_name) in enumerate(planning_documents.items()):
    doc_uploader(doc_id, doc_name, "PLANNING", cols[i % 2])

st.markdown("#### 2B — Standards documents")
cols = st.columns(2)
for i, (doc_id, doc_name) in enumerate(standard_documents.items()):
    doc_uploader(doc_id, doc_name, "PLANNING", cols[i % 2])

st.divider()

# ── STEP 3 : M1 INPUTS ────────────────────────────────────────────────────────
st.markdown("### Étape 3 — Inputs M1 (3 documents)")
st.caption(
    "Documents source pour la génération des HLR. Chacun peut être "
    "uploadé indépendamment — par exemple juste le Hardware doc seul."
)

col1, col2, col3 = st.columns(3)
doc_uploader("PDS", "Product Design Specification", "M1_INPUT", col1)
doc_uploader("ARCH", "Architecture Système", "M1_INPUT", col2)
doc_uploader("HW", "Hardware Documentation", "M1_INPUT", col3)

st.divider()

# ── SUMMARY ───────────────────────────────────────────────────────────────────
doc_id_to_module = {d: "PLANNING" for d in planning_docs.keys()}
doc_id_to_module.update({"PDS": "M1_INPUT", "ARCH": "M1_INPUT", "HW": "M1_INPUT"})
all_doc_ids = list(doc_id_to_module.keys())
saved = [d for d in all_doc_ids if get_artifact_by_module(active["id"], doc_id_to_module[d], d)]
st.markdown(f"### {len(saved)}/{len(all_doc_ids)} documents sauvegardés")
if len(saved) < len(all_doc_ids):
    missing = [d for d in all_doc_ids if d not in saved]
    st.info(f"Documents manquants : {', '.join(missing)}. Le pipeline peut déjà tourner avec ce qui est présent — les modules utiliseront ce qui est disponible en base.")
else:
    st.success("Tous les documents sont présents.")
    st.page_link("pages/2_pipeline.py", label="Aller au Pipeline")